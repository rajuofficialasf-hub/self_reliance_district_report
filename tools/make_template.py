"""টেমপ্লেট (.xlsx) বানায়। কলামের নাম নেওয়া হয় src/app.js-এর KEYS তালিকা থেকে।
চালাতে: pip install openpyxl && python3 tools/make_template.py
"""
import re, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
import pathlib
ROOT=pathlib.Path(__file__).resolve().parent.parent
js=(ROOT/'src'/'app.js').read_text(encoding='utf-8')
blk=js[js.index('var KEYS'):js.index('var REQUIRED')]
KEYS={m.group(1):re.findall(r"'([^']+)'",m.group(2))[0] for m in re.finditer(r"(\w+):\[([^\]]+)\]",blk)}
REQ=['name','dist','up','gender','plan','hhInc','debt']
YN=['হ্যাঁ','না']
LISTS={
 'gender':['পুরুষ','মহিলা'],
 'marital':['বিবাহিত','অবিবাহিত','বিধবা / বিপত্নীক','তালাকপ্রাপ্ত'],
 'edu':['নিরক্ষর','প্রাথমিক অসম্পূর্ণ','প্রাথমিক সম্পূর্ণ','মাধ্যমিক অসম্পূর্ণ','SSC সমমান','দাখিল','HSC সমমান','আলিম',"স্নাতক (Bachelor's) সমমান",'দাওরায়ে হাদিস (তাকমিল)','হিফজুল কুরআন'],
 'occ':['দিনমজুর','ব্যবসা','গৃহকর্মী','কৃষক','বেকার','রিক্সা/ভ্যান চালক','মৎস্যজীবী','শিক্ষার্থী','অন্যান্য'],
 'multi':YN,'live':YN,'allow':YN,'hasLoan':YN,'loanMet':YN,
 'plan':['ব্যবসা করে','গবাদিপশু লালনপালন করে'],
 'animal':['গরু','ছাগল'],
 'house':['টিনশেড ঘর','কাঁচা ঘর','সেমিপাকা ঘর','বাঁশের ঘর','ছনের ঘর','ভাড়া বাসা','গৃহহীন','অন্যান্য'],
 'loanType':['কর্জে হাসানা','এনজিও','ব্যাংক লোন','মহাজনী','অন্যান্য'],
 'loanPur':['দৈনন্দিন চাহিদা পূরণ','চিকিৎসা ব্যয়','মেয়ের বিয়ের জন্য','ব্যবসা','ঋণ পরিশোধ','ঘর নির্মাণ / মেরামত','শিক্ষা ব্যয়','কৃষি'],
 'bizType':['ফল/শাক-সবজি/বাদাম','মুদি দোকান','রিকশা/ভ্যান/অটোমিশুক','চা-কফি/জুস/আইসক্রিম','হাঁস মুরগি পালন','ফেরিওয়ালা','আতর/টুপি/তাসবীহ','মাছ/দুধ/শুটকি','পোল্ট্রি খামার','মাছের ঘের','কাপড়','কৃষি/চাষাবাদ','কসমেটিকস/মেয়েদের পণ্য','বাঁশ-বেত/ঝুড়ি','জুতা/মুচি','মোবাইল এক্সেসরিজ/ঘড়ি','পিঠা/ডিম/পপকর্ণ/চিপস','আচার/গুড়/চিড়া/নারিকেল','প্লাস্টিক/স্টিল পণ্য','কাঠ/শোপিস','অন্যান্য'],
}
ORDER=['com','imam','elig','appr','name','gender','marital','edu','dob','div','dist','up','uni','vil','kids','mem','occ','ownInc','exp','multi','hhInc','plan','bizType','need','animal','cap','house','live','allow','hasLoan','loanType','loanPur','loanMet','loanAmt','borrow','dower','debt','net']
assert set(ORDER)==set(KEYS), set(KEYS)^set(ORDER)
DESC={
 'com':'কমিউনিটির নম্বর, যেমন ASF-CBSRP-2651। প্রতিটি কমিউনিটির জন্য আলাদা গণনা হয়।',
 'imam':'কমিউনিটির স্থানীয় ইমামের নাম। কমিউনিটি তালিকায় দেখানো হয়।',
 'elig':'আলেমের অনুমোদন হলে “অনুমোদিত” লিখুন। সব সারিতে থাকলে রিপোর্টে নোট আসে।',
 'appr':'অনুমোদনের তারিখ। Excel তারিখ বা 2026-06-15 অথবা 15/06/2026 লিখুন। সময়রেখা চার্টে লাগে।',
 'name':'উপকারভোগীর পুরো নাম। এই ঘর ফাঁকা সারি গণনায় ধরা হয় না।',
 'gender':'ড্রপডাউন থেকে বাছুন: পুরুষ বা মহিলা।',
 'marital':'ড্রপডাউন থেকে বাছুন।',
 'edu':'ড্রপডাউন থেকে সর্বোচ্চ শিক্ষাগত যোগ্যতা বাছুন।',
 'dob':'জন্মতারিখ। Excel তারিখ বা 1975-05-10 অথবা 10/05/1975। বয়সের চার্টে লাগে।',
 'div':'বিভাগের নাম, যেমন Rangpur (রংপুর)।',
 'dist':'জেলার নাম ইংরেজিতে, যেমন Gaibandha। একই ফাইলে একাধিক জেলা থাকলে টুলে জেলা বাছা যায়।',
 'up':'উপজেলার নাম ইংরেজিতে, যেমন Sundarganj। বাংলায় লিখলেও চলবে।',
 'uni':'ইউনিয়নের নাম। শুধু ইউনিয়ন সংখ্যা গণনায় লাগে।',
 'vil':'গ্রামের নাম। রিপোর্টে সরাসরি দেখানো হয় না।',
 'kids':'পরিবারে ১৪ বছরের কম বয়সী শিশুর সংখ্যা (সংখ্যায়)।',
 'mem':'পরিবারের মোট সদস্য সংখ্যা (সংখ্যায়)। মাথাপিছু আয় হিসাবে লাগে।',
 'occ':'ড্রপডাউন থেকে প্রধান পেশা বাছুন।',
 'ownInc':'উপকারভোগীর নিজের মাসিক আয়, টাকায়।',
 'exp':'উপকারভোগীর মাসিক ব্যয়, টাকায়।',
 'multi':'পরিবারে একাধিক আয়কারী থাকলে হ্যাঁ, নইলে না।',
 'hhInc':'পরিবারের মোট মাসিক আয়, টাকায় (সংখ্যায়)। আয়ের ব্যান্ড ও মাথাপিছু আয় এখান থেকে আসে।',
 'plan':'ড্রপডাউন থেকে বাছুন: ব্যবসা করে বা গবাদিপশু লালনপালন করে।',
 'bizType':'ব্যবসা বেছে থাকলে কী ধরনের ব্যবসা, তালিকা থেকে বাছুন।',
 'need':'ব্যবসার জন্য কত টাকা লাগবে, টাকায় (সংখ্যায়)। শুধু ব্যবসা বেছে থাকলে।',
 'animal':'গবাদিপশু বেছে থাকলে গরু বা ছাগল।',
 'cap':'বর্তমানে নিজের কাছে কত টাকা পুঁজি আছে, টাকায়। না থাকলে 0।',
 'house':'ড্রপডাউন থেকে ঘরের ধরন বাছুন।',
 'live':'বর্তমানে গবাদিপশু থাকলে হ্যাঁ, নইলে না।',
 'allow':'প্রতিবন্ধী, বয়স্ক বা বিধবা ভাতা পেলে হ্যাঁ, নইলে না।',
 'hasLoan':'কোনো ঋণ বা কিস্তি থাকলে হ্যাঁ, নইলে না।',
 'loanType':'ঋণ থাকলে উৎস বাছুন।',
 'loanPur':'ঋণ থাকলে ঋণ নেওয়ার কারণ বাছুন।',
 'loanMet':'ঋণের উদ্দেশ্য পূরণ হয়ে থাকলে হ্যাঁ, নইলে না।',
 'loanAmt':'ঋণ থাকলে মোট ঋণের পরিমাণ, টাকায়।',
 'borrow':'চিকিৎসা, সাংসারিক বা একান্ত প্রয়োজনে নেওয়া ঋণের পরিমাণ, টাকায়। না থাকলে 0।',
 'dower':'অনাদায়ী দেনমোহরের পরিমাণ, টাকায়। না থাকলে 0।',
 'debt':'মোট দেনা, টাকায়। হিসাব: নেওয়া ঋণ + অনাদায়ী দেনমোহর + বকেয়া বিল ও খরচ। দেনমোহর ও ঋণের ভাগ এই যোগফল থেকে বের হয়।',
 'net':'নিট সম্পদ, টাকায় (মোট সম্পদ বিয়োগ মোট দেনা)। ঋণাত্মক হতে পারে।',
}
wb=openpyxl.Workbook()
ws=wb.active; ws.title='ডাটা'
gi=wb.create_sheet('নির্দেশনা'); ls=wb.create_sheet('ড্রপডাউন তালিকা')
thin=Side(style='thin',color='C9CED3')
req_fill=PatternFill('solid',fgColor='1F5F4B'); opt_fill=PatternFill('solid',fgColor='5B6670')
for i,k in enumerate(ORDER,1):
    c=ws.cell(1,i,KEYS[k]); c.font=Font(bold=True,color='FFFFFF'); c.fill=req_fill if k in REQ else opt_fill
    c.alignment=Alignment(wrap_text=True,vertical='center'); ws.column_dimensions[L(i)].width=22 if k not in ('debt','borrow') else 26
ws.row_dimensions[1].height=64; ws.freeze_panes='F2' if False else 'A2'
# lists sheet
col=1; listcols={}
for k,v in LISTS.items():
    pass
done={}
for k in ORDER:
    if k in LISTS:
        key=tuple(LISTS[k])
        if key not in done:
            c=len(done)+1; ls.cell(1,c,'তালিকা '+str(c)).font=Font(bold=True)
            for r,val in enumerate(LISTS[k],2): ls.cell(r,c,val)
            ls.column_dimensions[L(c)].width=30; done[key]=(c,len(LISTS[k]))
        c,n=done[key]
        dv=DataValidation(type='list',formula1=f"='ড্রপডাউন তালিকা'!${L(c)}$2:${L(c)}${n+1}",allow_blank=True,errorStyle='warning',showErrorMessage=True,errorTitle='তালিকার বাইরের মান',error='এই মানটি তালিকায় নেই। তবুও রাখতে চাইলে “হ্যাঁ” চাপুন।')
        ws.add_data_validation(dv); dv.add(f"{L(ORDER.index(k)+1)}2:{L(ORDER.index(k)+1)}1000")
for k in ('appr','dob'):
    ci=ORDER.index(k)+1
    for r in range(2,1001): ws.cell(r,ci).number_format='yyyy-mm-dd'
# guidance
gi.column_dimensions['A'].width=46; gi.column_dimensions['B'].width=16; gi.column_dimensions['C'].width=100
rows=[('SRP-২৬ ডাটা টেমপ্লেট: ব্যবহারের নির্দেশনা',),
('১. “ডাটা” শিটে প্রতিটি উপকারভোগীর জন্য একটি করে সারি পূরণ করুন। প্রথম সারির হেডার বদলাবেন না।',),
('২. গাঢ় সবুজ হেডারের ঘর বাধ্যতামূলক। ধূসর হেডারের ঘর ফাঁকা থাকলে রিপোর্টের সংশ্লিষ্ট অংশ ফাঁকা বা অসম্পূর্ণ দেখাবে।',),
('৩. যেসব ঘরে ড্রপডাউন আছে সেখানে তালিকা থেকে বাছুন। টাকার অঙ্ক সংখ্যায় লিখুন (কমা ছাড়া)।',),
('৪. ফাইলটি .xlsx হিসেবে সেভ করে রিপোর্ট জেনারেটরে ছেড়ে দিন। একটি ফাইলে একাধিক জেলা রাখা যায়।',),
('৫. Kobo থেকে নামানো ফাইলেও চলবে, কারণ টুল হেডারের নাম দেখে কলাম চেনে।',),
('',),('কলামের নাম','প্রয়োজন','কী লিখবেন')]
for r in rows: gi.append(list(r))
gi['A1'].font=Font(bold=True,size=14)
for c in gi[8]: c.font=Font(bold=True,color='FFFFFF'); c.fill=req_fill
for k in ORDER:
    gi.append([KEYS[k],'বাধ্যতামূলক' if k in REQ else 'ঐচ্ছিক',DESC[k]])
for row in gi.iter_rows(min_row=9):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
wb.active=0
(ROOT/'template').mkdir(exist_ok=True)
wb.save(str(ROOT/'template'/'SRP-26_Data_Template.xlsx'))
print(len(ORDER),'columns')
